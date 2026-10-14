"""Excel workbook builder for the daily report.

The workbook is assembled in two steps. :func:`tables` turns a :class:`BriefInputs` into
plain row lists (one list per sheet, header row first, only ``str``/``int``/``float``/
``bool``/``None`` values), so the Google Sheets publisher can reuse the same content.
:func:`build_workbook` writes those tables into an ``openpyxl`` workbook, applies number
formats, bold frozen headers, column widths and a colour scale on correlation sheets, and
returns the ``.xlsx`` bytes. Sheet order: ``Summary``, ``Pearson <w>`` / ``Spearman <w>``
per configured window that has matrices, ``Uncorrelated Sets``, ``Momentum``,
``Positions``, ``Data Quality``, ``Parameters``.
"""

from __future__ import annotations

import math
from collections import Counter
from collections.abc import Iterable
from dataclasses import fields
from datetime import date
from io import BytesIO

import numpy as np
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from uncorrelated_trading.models import (
    AnalysisResult,
    CorrelationMatrix,
    PortfolioContext,
    UncorrelatedSet,
)
from uncorrelated_trading.report.inputs import BriefInputs, holdings_rows

# ---------------------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------------------

MAX_SHEET_NAME = 31
MAX_COLUMN_WIDTH = 40
RHO_FORMAT = "0.000"
PCT_FORMAT = "0.00%"
MONEY_FORMAT = "#,##0.00"

SET_COLUMNS = [
    "size",
    "rank",
    "is_largest",
    "symbols",
    "max_abs_rho",
    "max_abs_rho_ci_upper",
    "mean_abs_rho",
    "min_eigenvalue",
    "spearman_max_abs_rho",
    "passes_spearman",
    "stable_in_long_window",
    "mean_momentum_z",
]
MOMENTUM_COLUMNS = [
    "rank",
    "symbol",
    "direction",
    "composite_z",
    "trailing_return",
    "skip_return",
    "vol_adjusted_return",
    "realized_vol",
]
HOLDINGS_COLUMNS = [
    "symbol",
    "qty",
    "current_price",
    "market_value",
    "unrealized_pl",
    "unrealized_plpc",
    "momentum_rank",
    "direction",
    "in_largest_set",
]
QUALITY_COLUMNS = ["symbol", "status", "detail", "bars_available", "source"]

# Column header -> number format, applied to tabular sheets.
COLUMN_FORMATS: dict[str, str] = {
    "max_abs_rho": RHO_FORMAT,
    "max_abs_rho_ci_upper": RHO_FORMAT,
    "mean_abs_rho": RHO_FORMAT,
    "min_eigenvalue": RHO_FORMAT,
    "spearman_max_abs_rho": RHO_FORMAT,
    "mean_momentum_z": RHO_FORMAT,
    "composite_z": RHO_FORMAT,
    "vol_adjusted_return": RHO_FORMAT,
    "trailing_return": PCT_FORMAT,
    "skip_return": PCT_FORMAT,
    "realized_vol": PCT_FORMAT,
    "unrealized_plpc": PCT_FORMAT,
    "current_price": MONEY_FORMAT,
    "market_value": MONEY_FORMAT,
    "unrealized_pl": MONEY_FORMAT,
}
# Key (first column) -> number format for the value cell, applied to key/value rows.
KEY_FORMATS: dict[str, str] = {
    "cash": MONEY_FORMAT,
    "equity": MONEY_FORMAT,
    "buying power": MONEY_FORMAT,
    "portfolio value": MONEY_FORMAT,
    "last equity": MONEY_FORMAT,
    "day P&L": MONEY_FORMAT,
    "account equity": MONEY_FORMAT,
    "account cash": MONEY_FORMAT,
    "account day P&L": MONEY_FORMAT,
}

Row = list[object]

# ---------------------------------------------------------------------------------------
# Value helpers
# ---------------------------------------------------------------------------------------


def _plain(value: object) -> object:
    """Convert numpy scalars to Python scalars and non-finite floats to ``None``."""
    if isinstance(value, np.generic):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if value is None or isinstance(value, str | bool | int | float):
        return value
    return str(value)


def _join(items: Iterable[object]) -> str:
    """Comma-join items as text."""
    return ", ".join(str(i) for i in items)


def _sheet_name(name: str) -> str:
    """Validate an Excel sheet name (1..31 characters)."""
    if not name or len(name) > MAX_SHEET_NAME:
        raise ValueError(f"invalid sheet name {name!r}: must be 1..{MAX_SHEET_NAME} characters")
    return name


# ---------------------------------------------------------------------------------------
# Table builders
# ---------------------------------------------------------------------------------------


def _summary_rows(inputs: BriefInputs) -> list[Row]:
    """Key/value rows describing the run and the account headline."""
    r = inputs.result
    largest = r.largest_set
    rows: list[Row] = [
        ["field", "value"],
        ["as-of day", r.as_of.day.isoformat()],
        ["mode", r.as_of.mode],
        ["provisional", r.as_of.provisional],
        ["calendar source", r.as_of.calendar_source],
        ["as-of note", r.as_of.note],
        ["universe size", len(r.universe)],
        ["analysed", len(r.quality.analysed)],
        ["run id", r.run_id],
        ["created at", r.created_at.isoformat()],
        ["tau", r.params.tau],
        ["windows", _join(r.params.windows)],
        ["largest set size", largest.size if largest else 0],
        ["largest set", _join(largest.symbols) if largest else "none"],
        ["largest set exact", r.largest_set_exact],
        ["app version", inputs.app_version],
    ]
    account = inputs.portfolio.account
    if account is not None:
        rows.append(["account equity", account.equity])
        rows.append(["account cash", account.cash])
        rows.append(["account day P&L", account.equity - account.last_equity])
    else:
        rows.append(["account error", inputs.portfolio.error or "account unavailable"])
    return rows


def _matrix_rows(matrix: CorrelationMatrix) -> list[Row]:
    """Square matrix with symbols as header row and first column."""
    rows: list[Row] = [["symbol", *matrix.symbols]]
    for i, sym in enumerate(matrix.symbols):
        rows.append([sym, *(float(v) for v in matrix.values[i])])
    return rows


def _set_row(s: UncorrelatedSet, rank: int, is_largest: bool) -> Row:
    """One row of the ``Uncorrelated Sets`` sheet."""
    return [
        s.size,
        rank,
        is_largest,
        _join(s.symbols),
        s.max_abs_rho,
        s.max_abs_rho_ci_upper,
        s.mean_abs_rho,
        s.min_eigenvalue,
        s.spearman_max_abs_rho,
        s.passes_spearman,
        s.stable_in_long_window,
        s.mean_momentum_z,
    ]


def _sets_rows(r: AnalysisResult) -> list[Row]:
    """Pairs, then sets by ascending size, then the largest set flagged ``is_largest``."""
    rows: list[Row] = [list(SET_COLUMNS)]
    rows.extend(_set_row(s, i, False) for i, s in enumerate(r.pairs, start=1))
    for size in range(3, r.params.max_set_size_exact + 1):
        sets = r.sets_by_size.get(size, ())
        rows.extend(_set_row(s, i, False) for i, s in enumerate(sets, start=1))
    if r.largest_set is not None:
        rows.append(_set_row(r.largest_set, 1, True))
    return rows


def _momentum_rows(r: AnalysisResult) -> list[Row]:
    """Momentum scores in rank order."""
    rows: list[Row] = [list(MOMENTUM_COLUMNS)]
    for m in r.momentum:
        rows.append(
            [
                m.rank,
                m.symbol,
                m.direction,
                m.composite_z,
                m.trailing_return,
                m.skip_return,
                m.vol_adjusted_return,
                m.realized_vol,
            ]
        )
    return rows


def _account_rows(portfolio: PortfolioContext) -> list[Row]:
    """Account key/value rows, or the error text when the account could not be read."""
    a = portfolio.account
    if a is None:
        return [["account error", portfolio.error or "account unavailable"]]
    return [
        ["cash", a.cash],
        ["equity", a.equity],
        ["buying power", a.buying_power],
        ["portfolio value", a.portfolio_value],
        ["last equity", a.last_equity],
        ["day P&L", a.equity - a.last_equity],
        ["currency", a.currency],
        ["paper", a.paper],
    ]


def _positions_rows(inputs: BriefInputs) -> list[Row]:
    """Join holdings with momentum and set membership, then a blank row and the account."""
    rows: list[Row] = [list(HOLDINGS_COLUMNS)]
    rows.extend([h[c] for c in HOLDINGS_COLUMNS] for h in holdings_rows(inputs))
    rows.append([])
    rows.extend(_account_rows(inputs.portfolio))
    return rows


def _quality_rows(r: AnalysisResult) -> list[Row]:
    """Per-symbol data-quality rows, a blank row, then counts."""
    q = r.quality
    rows: list[Row] = [list(QUALITY_COLUMNS)]
    rows.extend([s.symbol, s.status, s.detail, s.bars_available, s.source] for s in q.per_symbol)
    rows.append([])
    rows.append(["requested", len(q.requested)])
    rows.append(["analysed", len(q.analysed)])
    rows.append(["primary source failures", q.primary_source_failures])
    rows.append(["fallback calls", q.fallback_calls])
    for status, count in sorted(Counter(s.status for s in q.per_symbol).items()):
        rows.append([f"status {status}", count])
    return rows


def _parameter_rows(r: AnalysisResult) -> list[Row]:
    """One row per ``AnalysisParams`` field; tuples are comma-joined."""
    rows: list[Row] = [["parameter", "value"]]
    for f in fields(r.params):
        value = getattr(r.params, f.name)
        rows.append([f.name, _join(value) if isinstance(value, tuple) else value])
    return rows


def _correlation_sheets(r: AnalysisResult) -> dict[str, list[Row]]:
    """``Pearson <w>`` and ``Spearman <w>`` tables for each window that has matrices."""
    out: dict[str, list[Row]] = {}
    for w in r.params.windows:
        for method, label in (("pearson", "Pearson"), ("spearman", "Spearman")):
            matrix = next((c for c in r.correlations if c.window == w and c.method == method), None)
            if matrix is not None:
                out[f"{label} {w}"] = _matrix_rows(matrix)
    return out


def tables(inputs: BriefInputs) -> dict[str, list[list[object]]]:
    """Build every report table as plain rows.

    Parameters
    ----------
    inputs : BriefInputs
        Report inputs.

    Returns
    -------
    dict[str, list[list[object]]]
        Sheet name -> rows, header row first, in workbook order. Values are plain Python
        ``str``/``int``/``float``/``bool``/``None`` (numpy scalars converted, non-finite
        floats replaced by ``None``). Empty lists are blank separator rows.

    Raises
    ------
    ValueError
        If a sheet name is empty or longer than 31 characters.
    """
    r = inputs.result
    raw: dict[str, list[Row]] = {"Summary": _summary_rows(inputs)}
    raw.update(_correlation_sheets(r))
    raw["Uncorrelated Sets"] = _sets_rows(r)
    raw["Momentum"] = _momentum_rows(r)
    raw["Positions"] = _positions_rows(inputs)
    raw["Data Quality"] = _quality_rows(r)
    raw["Parameters"] = _parameter_rows(r)
    return {
        _sheet_name(name): [[_plain(v) for v in row] for row in rows] for name, rows in raw.items()
    }


# ---------------------------------------------------------------------------------------
# Workbook writer
# ---------------------------------------------------------------------------------------


def _write_sheet(ws: Worksheet, rows: list[list[object]], is_matrix: bool) -> None:
    """Write rows into ``ws`` and apply header, number formats, panes and widths."""
    for row in rows:
        ws.append(row)
    if not rows:
        return
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "B2" if is_matrix else "A2"
    header = rows[0]
    if is_matrix:
        n = len(header) - 1
        for r_idx in range(2, n + 2):
            for c_idx in range(2, n + 2):
                ws.cell(row=r_idx, column=c_idx).number_format = RHO_FORMAT
        if n > 0:
            ws.conditional_formatting.add(
                f"B2:{get_column_letter(n + 1)}{n + 1}",
                ColorScaleRule(  # type: ignore[no-untyped-call]
                    start_type="num",
                    start_value=-1,
                    start_color="F8696B",
                    mid_type="num",
                    mid_value=0,
                    mid_color="FFFFFF",
                    end_type="num",
                    end_value=1,
                    end_color="5A8AC6",
                ),
            )
    else:
        formats = {i: COLUMN_FORMATS[str(h)] for i, h in enumerate(header) if h in COLUMN_FORMATS}
        tabular_end = len(rows)
        for r_idx, row in enumerate(rows[1:], start=2):
            if not row:
                tabular_end = r_idx
                break
            for c_idx, fmt in formats.items():
                ws.cell(row=r_idx, column=c_idx + 1).number_format = fmt
        for r_idx, row in enumerate(rows, start=1):
            if (r_idx == 1 or r_idx >= tabular_end) and len(row) >= 2 and row[0] in KEY_FORMATS:
                ws.cell(row=r_idx, column=2).number_format = KEY_FORMATS[str(row[0])]
    widths: dict[int, int] = {}
    for row in rows:
        for c_idx, value in enumerate(row, start=1):
            length = len(str(value)) if value is not None else 0
            widths[c_idx] = max(widths.get(c_idx, 0), length)
    for c_idx, length in widths.items():
        ws.column_dimensions[get_column_letter(c_idx)].width = min(length + 2, MAX_COLUMN_WIDTH)


def build_workbook(inputs: BriefInputs) -> bytes:
    """Build the report workbook as ``.xlsx`` bytes.

    Parameters
    ----------
    inputs : BriefInputs
        Report inputs.

    Returns
    -------
    bytes
        The serialised workbook, one sheet per entry of :func:`tables` in that order.

    Raises
    ------
    ValueError
        If a sheet name is empty or longer than 31 characters.
    """
    wb = Workbook()
    default = wb.active
    if default is not None:
        wb.remove(default)
    for name, rows in tables(inputs).items():
        ws = wb.create_sheet(title=name)
        _write_sheet(ws, rows, is_matrix=name.startswith(("Pearson ", "Spearman ")))
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def workbook_filename(as_of_day: date) -> str:
    """Return the attachment filename for an as-of day.

    Parameters
    ----------
    as_of_day : date
        As-of trading day of the run.

    Returns
    -------
    str
        ``uncorrelated_YYYY-MM-DD.xlsx``.
    """
    return f"uncorrelated_{as_of_day.isoformat()}.xlsx"
