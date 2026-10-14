from dataclasses import replace
from datetime import date
from io import BytesIO

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from factories import sample_brief_inputs, sample_params
from uncorrelated_trading.models import AnalysisResult, PortfolioContext
from uncorrelated_trading.report.inputs import BriefInputs
from uncorrelated_trading.report.workbook import build_workbook, tables, workbook_filename


def _three_window_inputs() -> BriefInputs:
    base = sample_brief_inputs()
    r = base.result
    extra = tuple(replace(c, window=250, n_obs=250) for c in r.correlations if c.window == 60)
    result = replace(
        r,
        params=replace(r.params, windows=(60, 120, 250)),
        correlations=r.correlations + extra,
    )
    return replace(base, result=result)


def _load(inputs: BriefInputs) -> openpyxl.Workbook:
    return openpyxl.load_workbook(BytesIO(build_workbook(inputs)))


def _rows(ws: Worksheet) -> list[tuple[object, ...]]:
    return [tuple(row) for row in ws.iter_rows(values_only=True)]


def test_workbook_has_expected_sheets():
    wb = _load(_three_window_inputs())
    assert wb.sheetnames == [
        "Summary",
        "Pearson 60",
        "Spearman 60",
        "Pearson 120",
        "Spearman 120",
        "Pearson 250",
        "Spearman 250",
        "Uncorrelated Sets",
        "Momentum",
        "Positions",
        "Data Quality",
        "Parameters",
    ]
    assert all(len(name) <= 31 for name in wb.sheetnames)


def test_window_without_matrices_is_skipped():
    base = sample_brief_inputs()
    result = replace(base.result, params=replace(sample_params(), windows=(60, 120, 250)))
    names = list(tables(replace(base, result=result)))
    assert "Pearson 250" not in names and "Pearson 120" in names


def test_correlation_sheet_is_square_with_symbol_headers():
    inputs = sample_brief_inputs()
    wb = _load(inputs)
    ws = wb["Pearson 60"]
    rows = _rows(ws)
    symbols = ("AAA", "BBB", "CCC", "DDD", "EEE")
    assert ws.max_row == ws.max_column == len(symbols) + 1
    assert rows[0][1:] == symbols
    assert tuple(r[0] for r in rows[1:]) == symbols
    matrix = next(c for c in inputs.result.correlations if c.window == 60 and c.method == "pearson")
    for i in range(len(symbols)):
        for j in range(len(symbols)):
            assert rows[i + 1][j + 1] == float(matrix.values[i, j])
    assert ws.freeze_panes == "B2"
    assert ws["B2"].number_format == "0.000"
    assert ws["A1"].font.b and ws["B1"].font.b
    rules = list(ws.conditional_formatting)
    assert len(rules) == 1
    assert str(rules[0].sqref) == "B2:F6"
    scale = rules[0].rules[0].colorScale
    assert [float(v.val) for v in scale.cfvo] == [-1.0, 0.0, 1.0]
    assert [c.rgb[-6:] for c in scale.color] == ["F8696B", "FFFFFF", "5A8AC6"]


def test_momentum_sheet_rows_match_result():
    inputs = sample_brief_inputs()
    ws = _load(inputs)["Momentum"]
    rows = _rows(ws)
    assert rows[0] == (
        "rank",
        "symbol",
        "direction",
        "composite_z",
        "trailing_return",
        "skip_return",
        "vol_adjusted_return",
        "realized_vol",
    )
    assert len(rows) - 1 == len(inputs.result.momentum)
    for row, m in zip(rows[1:], inputs.result.momentum, strict=True):
        assert row == (
            m.rank,
            m.symbol,
            m.direction,
            m.composite_z,
            m.trailing_return,
            m.skip_return,
            m.vol_adjusted_return,
            m.realized_vol,
        )
    assert ws.freeze_panes == "A2"
    assert ws["E2"].number_format == "0.00%"
    assert ws["A1"].font.b


def test_positions_sheet_flags_largest_set_member():
    rows = _rows(_load(sample_brief_inputs())["Positions"])
    header = rows[0]
    assert header[0] == "symbol" and "in_largest_set" in header
    flag = header.index("in_largest_set")
    by_symbol = {r[0]: r for r in rows[1:3]}
    assert by_symbol["AAA"][flag] is True
    assert by_symbol["XYZ"][flag] is False
    assert all(v is None for v in rows[3])
    account = {r[0]: r[1] for r in rows[4:]}
    assert account["equity"] == 100_000.0
    assert account["day P&L"] == 1_000.0
    assert account["paper"] is True


def test_positions_sheet_shows_portfolio_error():
    base = sample_brief_inputs()
    inputs = replace(
        base, portfolio=PortfolioContext(account=None, positions=(), error="401 unauthorized")
    )
    rows = _rows(_load(inputs)["Positions"])
    assert any("401 unauthorized" in str(v) for r in rows for v in r)


def test_uncorrelated_sets_order_and_largest_flag():
    t = tables(sample_brief_inputs())["Uncorrelated Sets"]
    header = t[0]
    assert header[:4] == ["size", "rank", "is_largest", "symbols"]
    body = t[1:]
    assert [r[0] for r in body] == [2, 2, 2, 3, 3]
    assert body[0][3] == "AAA, DDD"
    assert [r[2] for r in body] == [False, False, False, False, True]
    assert body[-1][3] == "AAA, BBB, DDD"


def test_missing_set_size_key_is_empty():
    base = sample_brief_inputs()
    result: AnalysisResult = replace(base.result, sets_by_size={}, largest_set=None)
    t = tables(replace(base, result=result))["Uncorrelated Sets"]
    assert [r[0] for r in t[1:]] == [2, 2, 2]


def test_tables_contain_plain_python_types_only():
    for name, rows in tables(_three_window_inputs()).items():
        assert len(name) <= 31
        for row in rows:
            for v in row:
                assert v is None or type(v) in (str, int, float, bool), (name, v, type(v))


def test_summary_and_parameters_sheets():
    t = tables(sample_brief_inputs())
    summary = {r[0]: r[1] for r in t["Summary"][1:]}
    assert summary["as-of day"] == "2026-10-13"
    assert summary["run id"] == "0123456789abcdef0123456789abcdef"
    assert summary["largest set size"] == 3
    assert summary["windows"] == "60, 120"
    params = {r[0]: r[1] for r in t["Parameters"][1:]}
    assert params["tau"] == 0.2
    assert params["momentum_weights"] == "0.4, 0.3, 0.3"
    assert len(params) == 12
    dq = t["Data Quality"]
    assert dq[0][:2] == ["symbol", "status"]
    assert any(r[0] == "ZZZ" and r[1] == "dropped_missing" for r in dq)


def test_column_widths_capped():
    base = sample_brief_inputs()
    ws = _load(base)["Data Quality"]
    widths = [d.width for d in ws.column_dimensions.values()]
    assert widths and max(widths) <= 40


def test_filename():
    assert workbook_filename(date(2026, 10, 13)) == "uncorrelated_2026-10-13.xlsx"


def test_number_formats_on_positions_and_sets():
    wb = _load(sample_brief_inputs())
    pos = wb["Positions"]
    header = [c.value for c in pos[1]]
    assert pos.cell(row=2, column=header.index("market_value") + 1).number_format == "#,##0.00"
    assert pos.cell(row=2, column=header.index("unrealized_plpc") + 1).number_format == "0.00%"
    equity_row = next(r for r in range(1, pos.max_row + 1) if pos.cell(r, 1).value == "equity")
    assert pos.cell(equity_row, 2).number_format == "#,##0.00"
    assert pos.freeze_panes == "A2"
    sets = wb["Uncorrelated Sets"]
    sheader = [c.value for c in sets[1]]
    assert sets.cell(row=2, column=sheader.index("min_eigenvalue") + 1).number_format == "0.000"
