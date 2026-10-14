from dataclasses import replace
from io import BytesIO

import numpy as np
import openpyxl
import pytest
from hypothesis import given, settings
from hypothesis import strategies as st

from factories import sample_analysis_result, sample_brief_inputs
from uncorrelated_trading.models import AnalysisResult
from uncorrelated_trading.report.inputs import BriefInputs
from uncorrelated_trading.report.workbook import build_workbook, tables

ALL = ("AAA", "BBB", "CCC", "DDD", "EEE")


@st.composite
def small_inputs(draw: st.DrawFn) -> BriefInputs:
    symbols = tuple(
        sorted(draw(st.lists(st.sampled_from(ALL), min_size=3, max_size=5, unique=True)))
    )
    windows = tuple(
        sorted(draw(st.lists(st.sampled_from((60, 120)), min_size=1, max_size=2, unique=True)))
    )
    seed = draw(st.integers(0, 2**16))
    base = sample_analysis_result()
    rng = np.random.default_rng(seed)
    n = len(symbols)
    corrs = []
    for c in base.correlations:
        if c.window not in windows:
            continue
        m = rng.uniform(-1.0, 1.0, size=(n, n))
        m = (m + m.T) / 2.0
        np.fill_diagonal(m, 1.0)
        corrs.append(replace(c, symbols=symbols, values=m))
    keep = set(symbols)
    momentum = tuple(
        replace(m, rank=i)
        for i, m in enumerate((m for m in base.momentum if m.symbol in keep), start=1)
    )
    largest = base.largest_set
    if largest is not None and not set(largest.symbols) <= keep:
        largest = None
    result: AnalysisResult = replace(
        base,
        params=replace(base.params, windows=windows),
        correlations=tuple(corrs),
        pairs=tuple(p for p in base.pairs if set(p.symbols) <= keep),
        sets_by_size={
            k: tuple(s for s in v if set(s.symbols) <= keep) for k, v in base.sets_by_size.items()
        },
        largest_set=largest,
        momentum=momentum,
        quality=replace(base.quality, analysed=symbols),
    )
    return replace(sample_brief_inputs(), result=result)


@settings(max_examples=25, deadline=None)
@given(small_inputs())
def test_workbook_round_trip(inputs: BriefInputs) -> None:
    wb = openpyxl.load_workbook(BytesIO(build_workbook(inputs)))
    t = tables(inputs)
    assert wb.sheetnames == list(t)
    n = len(inputs.result.quality.analysed)
    assert wb["Momentum"].max_row - 1 == n
    for w in inputs.result.params.windows:
        for label in ("Pearson", "Spearman"):
            ws = wb[f"{label} {w}"]
            assert ws.max_row == ws.max_column == n + 1
    for name, rows in t.items():
        ws = wb[name]
        for r_idx, row in enumerate(rows, start=1):
            for c_idx, value in enumerate(row, start=1):
                # openpyxl stores "" as an empty cell
                expected = None if value == "" else value
                got = ws.cell(row=r_idx, column=c_idx).value
                if isinstance(expected, float):
                    # openpyxl serialises floats with ~16 significant digits
                    assert got == pytest.approx(expected, rel=1e-14, abs=1e-15)
                else:
                    assert got == expected
