from __future__ import annotations

import email
import io
import os
from email.message import Message
from pathlib import Path

import pytest
from openpyxl import load_workbook

from uncorrelated_trading.cli import main
from uncorrelated_trading.store.db import connect
from uncorrelated_trading.store.runs import RunStore

REPO = Path(__file__).resolve().parents[2]
OFFLINE_TOML = REPO / "tests" / "fixtures" / "config" / "offline.toml"
CSV = REPO / "tests" / "fixtures" / "csv" / "bars_small.csv"


def test_daily_dry_run_offline_end_to_end(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    for key in list(os.environ):
        monkeypatch.delenv(key)
    monkeypatch.chdir(tmp_path)
    data_dir = tmp_path / "data"

    code = main(
        [
            "--config",
            str(OFFLINE_TOML),
            "--data-dir",
            str(data_dir),
            "daily",
            "--offline",
            "--csv",
            str(CSV),
            "--dry-run",
            "--as-of",
            "2026-10-13",
        ]
    )

    assert code == 0
    emls = list((data_dir / "outbox").glob("*.eml"))
    assert len(emls) == 1
    msg = email.message_from_bytes(emls[0].read_bytes())
    assert "2026-10-13" in str(msg["Subject"])
    parts: list[Message] = [p for p in msg.walk() if p.get_filename()]
    xlsx = [p for p in parts if str(p.get_filename()).endswith(".xlsx")]
    assert len(xlsx) == 1
    payload = xlsx[0].get_payload(decode=True)
    assert isinstance(payload, bytes)
    wb = load_workbook(io.BytesIO(payload))
    assert "Momentum" in wb.sheetnames
    assert wb["Momentum"].max_row > 1

    db = data_dir / "uncorrelated_trading.sqlite3"
    assert db.is_file()
    conn = connect(db)
    try:
        run_id = RunStore(conn).latest_run_id()
        assert run_id is not None
    finally:
        conn.close()
